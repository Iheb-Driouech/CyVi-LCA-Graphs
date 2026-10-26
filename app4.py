# Import the necessary libraries
import streamlit as st
import pandas as pd
import numpy as np
import matplotlib
matplotlib.use("Agg")  # rendu sans affichage (serveur)
import matplotlib.pyplot as plt
from matplotlib import colors as mcolors
import os
import io
import re
import json
import zipfile
import hashlib
import threading


###############################
#   Constants
###############################

# Palette Okabe-Ito : distinguable par les personnes daltoniennes.
# Le noir d'origine est remplace par un gris (les hachures noires restent visibles).
OKABE_ITO = ["#0072B2", "#E69F00", "#009E73", "#D55E00",
             "#56B4E9", "#CC79A7", "#F0E442", "#999999"]

FONT_FAMILIES = ["DejaVu Sans", "sans-serif", "serif"]
FONT_WEIGHTS = ["normal", "bold"]
FONT_STYLES = ["normal", "italic"]
# (nom, libelle, min, max, defaut)
FONT_SIZE_SLIDERS = [
    ("title_size", "Title Font Size", 10, 30, 16),
    ("label_size", "Label Font Size", 8, 30, 12),
    ("text_size", "Text Font Size", 6, 30, 10),
    ("legend_size", "Legend Font Size", 6, 30, 10),
]
HATCH_CHOICES = ["///", "//", "/", "xx", "XX", "++", "--", "||", "..", "O", "oo", "None"]
DEFAULT_HATCHES = ["//", "oo", "xx", "--", "||", "++", "..", "None"]

PREVIEW_DPI = 110   # apercu a l'ecran (leger)
EXPORT_DPI = 300    # telechargement PNG
SETTINGS_VERSION = 1

# Les fonctions de trace utilisent des variables globales (couleurs, polices...).
# Streamlit sert plusieurs utilisateurs dans des threads differents : ce verrou
# empeche deux rendus simultanes de melanger leurs reglages.
_RENDER_LOCK = threading.Lock()


def palette_color(i):
    """Couleur n°i de la palette par defaut (Okabe-Ito)."""
    return OKABE_ITO[i % len(OKABE_ITO)]


class TemplateError(ValueError):
    """Le fichier Excel ne respecte pas la template : message lisible pour l'utilisateur."""


def unique_in_order(values):
    """Valeurs uniques dans l'ordre d'apparition (contrairement a sorted(set(...)))."""
    seen, out = set(), []
    for v in values:
        if v is None or (isinstance(v, float) and np.isnan(v)):
            continue
        if v not in seen:
            seen.add(v)
            out.append(v)
    return out


def _ss_default(key, value):
    """Initialise la valeur d'un widget dans st.session_state si elle n'existe pas encore."""
    if key not in st.session_state:
        st.session_state[key] = value






        
###############################
#   Function Definitions
###############################

# Reglages utilises par les fonctions de trace, remplis par apply_render_settings()
font_styles = None
scenario_colors = None
contributions_colors = None
scenario_hatches = {}

def get_font_styles():
    st.sidebar.subheader("🖋️ Font Style Settings")

    # Les widgets ont une cle fixe : leurs valeurs peuvent etre sauvegardees
    # et rechargees depuis un fichier de reglages JSON.
    _ss_default("font_family", FONT_FAMILIES[0])
    _ss_default("font_weight", FONT_WEIGHTS[0])
    _ss_default("font_style", FONT_STYLES[0])
    family = st.sidebar.selectbox("Font Family", FONT_FAMILIES, key="font_family")
    fontweight = st.sidebar.selectbox("Font Weight", FONT_WEIGHTS, key="font_weight")
    fontstyle = st.sidebar.selectbox("Font Style", FONT_STYLES, key="font_style")

    sizes = {}
    for name, label, lo, hi, default in FONT_SIZE_SLIDERS:
        _ss_default(f"font_{name}", default)
        sizes[name] = st.sidebar.slider(label, lo, hi, key=f"font_{name}")
    title_size, label_size = sizes["title_size"], sizes["label_size"]
    text_size, legend_size = sizes["text_size"], sizes["legend_size"]

    _ss_default("precision_mode_label", list(PRECISION_MODES.keys())[0])
    precision_label = st.sidebar.selectbox(
        "Totals: when 2-decimal rounding hides a difference",
        list(PRECISION_MODES.keys()), key="precision_mode_label",
        help="Totals use 2 decimals by default. If two different totals of the same impact "
             "category would look identical (e.g. 0.020412 and 0.016847 -> 0.02), this option "
             "chooses how more digits are shown. Example values in brackets."
    )

    # Les rcParams (police des indices/exposants) sont appliques au moment du rendu,
    # dans apply_render_settings(), sous verrou.
    return {
        "family": family,
        "weight": fontweight,
        "style": fontstyle,
        "title_size": title_size,
        "label_size": label_size,
        "text_size": text_size,
        "legend_size": legend_size,
        "precision_mode": PRECISION_MODES[precision_label]
    }


# Totaux de chaque categorie d'impact (tous scenarios), remplis dans main()
TOTAL_GROUPS = {}


def _fmt_total(v, extra=0, truncate=False):
    """Format de base : 2 decimales (ou notation scientifique a 2 decimales
    si |v| >= 1000 ou < 0.01), + `extra` decimales.
    truncate=True : on COUPE les chiffres au lieu d'arrondir."""
    from decimal import Decimal, ROUND_DOWN
    d = 2 + extra
    sci = v != 0 and (abs(v) >= 1000 or abs(v) < 0.01)
    if not truncate:
        return f"{v:.{d}E}" if sci else f"{v:.{d}f}"
    q = Decimal(1).scaleb(-d)  # 10^-d
    dv = Decimal(repr(float(v)))
    if sci:
        exp = dv.copy_abs().adjusted()
        mant = (dv.scaleb(-exp)).quantize(q, rounding=ROUND_DOWN)
        return f"{mant}E{exp:+03d}"
    return f"{dv.quantize(q, rounding=ROUND_DOWN)}"


# Modes proposes dans la sidebar quand l'arrondi a 2 decimales cache une difference
PRECISION_MODES = {
    "Truncate up to first differing digit + 1  (0.0204 / 0.0168)": "trunc_plus1",
    "Truncate, no rounding  (0.020 / 0.016)": "trunc",
    "Full value  (0.020412 / 0.016847)": "full",
}


def _fmt_full(v, max_sig=10):
    """Valeur complete (sans arrondi visible), zeros inutiles retires."""
    sci = v != 0 and (abs(v) >= 1000 or abs(v) < 0.01)
    if sci:
        mant, exp = f"{v:.{max_sig - 1}E}".split("E")
        mant = mant.rstrip("0").rstrip(".") if "." in mant else mant
        if "." not in mant:
            mant += ".0"
        return f"{mant}E{int(exp):+03d}"
    txt = f"{v:.{max_sig}g}"
    if "e" in txt.lower():
        txt = f"{v:.{max_sig}f}"
    if "." in txt:
        txt = txt.rstrip("0").rstrip(".")
    if "." not in txt or len(txt.split(".")[1]) < 2:   # au moins 2 decimales comme le format de base
        txt = f"{float(txt):.2f}"
    return txt


def _total_format_mode(group, mode="trunc_plus1", max_extra=6):
    """
    Renvoie (extra, truncate, full) pour une categorie d'impact :
    - (0, False, False) si l'arrondi a 2 decimales distingue deja les totaux differents ;
    - sinon, selon `mode` :
        "trunc_plus1" : troncature jusqu'au 1er chiffre different + 1  (0.0204 / 0.0168)
        "trunc"       : troncature jusqu'au 1er chiffre different      (0.020 / 0.016)
        "full"        : valeur complete                               (0.020412 / 0.016847)
    """
    try:
        vals = [float(x) for x in group if x is not None and np.isfinite(float(x))]
    except Exception:
        return 0, False, False

    def collide(extra, trunc):
        for i in range(len(vals)):
            for j in range(i + 1, len(vals)):
                if not np.isclose(vals[i], vals[j], rtol=1e-9, atol=0) and \
                        _fmt_total(vals[i], extra, trunc) == _fmt_total(vals[j], extra, trunc):
                    return True
        return False

    if not collide(0, False):
        return 0, False, False
    if mode == "full":
        return 0, False, True
    for extra in range(1, max_extra + 1):
        if not collide(extra, True):
            return (min(extra + 1, max_extra) if mode == "trunc_plus1" else extra), True, False
    return max_extra, True, False


def format_total(value, category=None):
    """
    Formate un total avec 2 decimales par defaut (arrondi). Si cet arrondi fait
    disparaitre la difference entre des scenarios de la meme categorie d'impact,
    on ajoute des decimales SANS arrondir (troncature) jusqu'au premier chiffre
    different + 1. Meme precision pour tous les
    scenarios d'une categorie, et dans tous les graphes.
    """
    try:
        v = float(value)
    except (TypeError, ValueError):
        return str(value)
    if np.isnan(v):
        return "nan"
    group = TOTAL_GROUPS.get(category) if category is not None else None
    mode = "trunc_plus1"
    try:
        if font_styles and font_styles.get("precision_mode"):
            mode = font_styles["precision_mode"]
    except Exception:
        pass
    extra, trunc, full = _total_format_mode(group, mode) if group is not None else (0, False, False)
    if full:
        return _fmt_full(v)
    return _fmt_total(v, extra, trunc)


def legend_hatch(hatch, target_lines=4, target_cross=3):
    """
    Hachure plus dense pour les LEGENDES : dans le petit rectangle de legende,
    une hachure "/" n'affiche qu'un seul trait. On augmente la densite :
    - traits simples (/ \\ | -) -> densite 4  (ex. "/" -> "////")
    - motifs croises (x X +)     -> densite 3  (ex. "xx" -> "xxx")
    - cercles et points (o O . *) inchanges (sinon pave noir).
    Le motif reste le meme que sur les barres.
    """
    if hatch in (None, "", "None"):
        return None
    out = ""
    for ch in dict.fromkeys(str(hatch)):          # caracteres distincts, ordre conserve
        k = str(hatch).count(ch)
        if ch in "oO.*":
            out += ch * k
        elif ch in "xX+":
            out += ch * max(k, target_cross)
        else:
            out += ch * max(k, target_lines)
    return out


def clean_scenario_name(name):
    """
    Extrait le nom entre parenthèses s'il y en a, sinon renvoie le nom complet.
    Ex: "Scenario 1 (Baseline)" -> "Baseline" ; "mechano" -> "mechano".
    """
    name = str(name)
    start = name.find("(")
    end = name.find(")")
    if start != -1 and end > start:
        return name[start + 1 : end].strip()
    return name.strip()


# Les indices/exposants sont rendus avec la police du texte (pas une police "maths")
plt.rcParams["mathtext.default"] = "regular"

_MATHTEXT_SPECIAL = {"\\": r"\backslash ", "{": r"\{", "}": r"\}", "$": r"\$", "%": r"\%",
                     "#": r"\#", "&": r"\&", "_": r"\_", "^": r"\wedge ", " ": r"\ "}


def _escape_plain(text):
    """Echappe le texte normal pour qu'il ne soit pas interprete par matplotlib."""
    return str(text).replace("$", r"\$")


def _to_mathtext(text, vert):
    """Convertit un morceau de texte en indice ('subscript') ou exposant ('superscript')."""
    esc = "".join(_MATHTEXT_SPECIAL.get(c, c) for c in str(text))
    return f"$_{{{esc}}}$" if vert == "subscript" else f"$^{{{esc}}}$"


def read_rich_category_labels(file_path, sheet_name, excel_rows):
    """
    Relit la colonne A (lignes Excel `excel_rows`, numerotees a partir de 1) avec openpyxl
    pour conserver les indices / exposants saisis dans Excel (ex. CO2 avec 2 en indice,
    m3 avec 3 en exposant) et les convertit en notation matplotlib (CO$_{2}$, m$^{3}$).
    Retourne None si la lecture echoue (on garde alors les libelles bruts).
    """
    try:
        from openpyxl import load_workbook
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        if hasattr(file_path, "seek"):
            file_path.seek(0)
        wb = load_workbook(file_path, rich_text=True, data_only=True)
        ws = wb.worksheets[sheet_name] if isinstance(sheet_name, int) else wb[sheet_name]
        labels = []
        for r in excel_rows:
            cell = ws.cell(row=r, column=1)
            value = cell.value
            if value is None:
                labels.append(None)
                continue
            if isinstance(value, CellRichText):
                parts = []
                for block in value:
                    if isinstance(block, TextBlock):
                        vert = getattr(block.font, "vertAlign", None) if block.font is not None else None
                        parts.append(_to_mathtext(block.text, vert) if vert in ("subscript", "superscript")
                                     else _escape_plain(block.text))
                    else:
                        parts.append(_escape_plain(block))
                labels.append("".join(parts))
            else:
                # Mise en forme appliquee a toute la cellule
                vert = getattr(cell.font, "vertAlign", None) if cell.font is not None else None
                labels.append(_to_mathtext(value, vert) if vert in ("subscript", "superscript")
                              else _escape_plain(value))
        return labels
    except Exception as e:
        print(f"Rich text labels not available, using plain text: {e}")
        return None
    finally:
        if hasattr(file_path, "seek"):
            file_path.seek(0)


def _cell_ref(row_idx, col_idx):
    """Reference Excel (ex. 'C5') a partir des index 0-based lus par pandas."""
    from openpyxl.utils import get_column_letter
    return f"{get_column_letter(int(col_idx) + 1)}{int(row_idx) + 1}"


def _list_cells(refs, max_shown=8):
    shown = ", ".join(refs[:max_shown])
    return shown + (f" … (+{len(refs) - max_shown} more)" if len(refs) > max_shown else "")


def analyze_excel_and_generate_tables(file_path, sheet_name=0):
    """
    Analyzes the specified Excel file and sheet, identifying scenarios, contributions,
    and generating initial and combined tables for impact assessment.

    - Scenarios, contributions and impact categories keep the order of the Excel file.
    - Raises TemplateError with a readable message when the sheet does not follow
      the template (missing names, duplicates, non-numeric values...).
    """
    data = pd.read_excel(file_path, header=None, sheet_name=sheet_name)
    print(f"Excel file loaded successfully from sheet: {sheet_name}!")

    # Colonnes et lignes totalement vides (mise en forme sans valeur) : ignorees.
    # Les index pandas = position Excel - 1 : on les garde pour les messages d'erreur.
    data = data.dropna(axis=1, how="all")
    header = data.iloc[:2]
    body = data.iloc[2:].dropna(axis=0, how="all")
    if data.shape[0] < 3 or data.shape[1] < 2 or body.empty:
        raise TemplateError(
            "The sheet is too small: expected scenarios in row 1, contributions in row 2 "
            "and impact categories in column A starting from row 3."
        )
    first_col = data.columns[0]
    value_cols = list(data.columns[1:])

    def _clean_text(v):
        return None if pd.isna(v) or str(v).strip() == "" else str(v).strip()

    scenario_row = pd.Series([_clean_text(v) for v in header.iloc[0, 1:]], index=value_cols)
    contribution_row = pd.Series([_clean_text(v) for v in header.iloc[1, 1:]], index=value_cols)

    # Step 1: names of scenarios / contributions must be filled in
    errors = []
    missing_scen = [_cell_ref(header.index[0], c) for c in value_cols if scenario_row[c] is None]
    missing_contrib = [_cell_ref(header.index[1], c) for c in value_cols if contribution_row[c] is None]
    if missing_scen:
        errors.append(f"Scenario name missing in row 1: {_list_cells(missing_scen)}.")
    if missing_contrib:
        errors.append(f"Contribution name missing in row 2: {_list_cells(missing_contrib)}.")

    categories_raw = [_clean_text(v) for v in body[first_col]]
    missing_cat = [_cell_ref(i, first_col) for i, v in zip(body.index, categories_raw) if v is None]
    if missing_cat:
        errors.append(f"Impact category name missing in column A: {_list_cells(missing_cat)}.")
    dup_cat = [c for c in unique_in_order(categories_raw) if c is not None and categories_raw.count(c) > 1]
    if dup_cat:
        errors.append("The same impact category appears several times: "
                      + ", ".join(f"'{c}'" for c in dup_cat) + ".")
    if errors:
        raise TemplateError(" ".join(errors))

    # Step 2: short scenario names (text in brackets) must be unique
    scenario_full_names = unique_in_order(scenario_row)
    short_to_full = {}
    for full in scenario_full_names:
        short_to_full.setdefault(clean_scenario_name(full).lower(), []).append(full)
    dup_short = {s: f for s, f in short_to_full.items() if len(f) > 1}
    if dup_short:
        details = "; ".join(f"'{s}' ← " + ", ".join(f"\"{n}\"" for n in f) for s, f in dup_short.items())
        raise TemplateError(
            "Several scenarios have the same short name (the text in brackets is used as the name, "
            f"case ignored): {details}. Please give each scenario a different name."
        )
    scenario_names_cleaned = list(short_to_full.keys())   # ordre du fichier Excel
    print(f"Detected scenarios ({len(scenario_names_cleaned)}): {scenario_names_cleaned}")

    # Step 3: a contribution must appear only once per scenario
    pairs = list(zip(scenario_row, contribution_row))
    dup_pairs = [p for p in unique_in_order(pairs) if pairs.count(p) > 1]
    if dup_pairs:
        details = "; ".join(f"'{c}' in \"{s}\"" for s, c in dup_pairs)
        raise TemplateError(f"A contribution appears twice in the same scenario: {details}.")
    print(f"Detected contributions: {unique_in_order(contribution_row)}")

    # Step 4: values must be numbers (empty cells = 0)
    raw_values = body[value_cols]
    values = raw_values.apply(pd.to_numeric, errors="coerce")
    bad = [_cell_ref(r, c) + f" ('{raw_values.loc[r, c]}')"
           for c in value_cols for r in raw_values.index
           if pd.isna(values.loc[r, c]) and not pd.isna(raw_values.loc[r, c])
           and str(raw_values.loc[r, c]).strip() != ""]
    if bad:
        raise TemplateError(f"Non-numeric values found: {_list_cells(bad)}. "
                            "Use numbers only (decimal point or comma as set in Excel), or leave the cell empty.")
    values = values.fillna(0.0).astype(float)

    # Step 5: structure the data
    # Conserver les indices / exposants saisis dans Excel (CO2, m3, ...)
    excel_rows = [int(i) + 1 for i in body.index]
    rich_labels = read_rich_category_labels(file_path, sheet_name, excel_rows)
    if rich_labels is not None and len(rich_labels) == len(categories_raw):
        categories = [rl if rl is not None else c for rl, c in zip(rich_labels, categories_raw)]
    else:
        categories = categories_raw
    values.columns = pd.MultiIndex.from_arrays([list(scenario_row), list(contribution_row)])
    values.index = pd.Index(categories)

    # Initial table
    initial_table = values.copy()

    # Impact totals (absolute values) — sort=False : ordre du fichier Excel conserve
    total_impact_table = values.abs().T.groupby(level=0, sort=False).sum().T

    # Replace columns with simplified scenario names
    total_impact_table.columns = [
        clean_scenario_name(name)
        for name in total_impact_table.columns
    ]

    # Relative percentages
    relative_percentage_table = total_impact_table.copy()
    for category in relative_percentage_table.index:
        max_value = relative_percentage_table.loc[category].max()
        relative_percentage_table.loc[category] = (
            relative_percentage_table.loc[category] / max_value
        ) * 100
    relative_percentage_table = relative_percentage_table.add_suffix(" (%)")

    # Combine tables
    combined_table = pd.concat([total_impact_table, relative_percentage_table], axis=1)

    return initial_table, combined_table, total_impact_table, scenario_names_cleaned


def extract_contributions_from_initial_table(initial_table):
    """
    Extrait les contributions uniques depuis les colonnes MultiIndex de initial_table,
    dans l'ordre du fichier Excel.
    """
    try:
        if not isinstance(initial_table.columns, pd.MultiIndex):
            raise ValueError("initial_table n'a pas un MultiIndex en colonnes.")

        contributions = initial_table.columns.get_level_values(1)
        unique_contributions = unique_in_order(contributions.astype(str).str.strip())

        print(f"Contributions uniques détectées ({len(unique_contributions)}) : {unique_contributions}")
        return unique_contributions

    except Exception as e:
        print(f"Erreur lors de l'extraction depuis initial_table : {e}")
        return []

def generate_percentage_table(initial_table, total_impact_table):
    """
    Generates a DataFrame identical to initial_table, but converts its values to percentages
    relative to the maximum impact found in total_impact_table for each category.
    """
    try:
        percentage_table = initial_table.copy()

        for category in initial_table.index:
            # Get the maximum impact for this category
            max_total_impact = total_impact_table.loc[category].max()
            # Calculate percentages
            percentage_table.loc[category] = (
                initial_table.loc[category] / max_total_impact
            ) * 100

        return percentage_table
    except Exception as e:
        print(f"Error generating percentage table: {e}")
        return None


def generate_tables_by_scenario(initial_table, scenario_names_cleaned):
    """
    Creates a dictionary of tables, one per scenario, containing:
    - Impact categories
    - Contributions
    - Total impact per category
    - Percentage of each contribution relative to the total impact
    """
    try:
        scenario_tables = {}
        # Nom court (texte entre parentheses, en minuscules) -> nom complet de la colonne.
        # Correspondance EXACTE : "lab1" ne doit pas recuperer les donnees de "lab10".
        full_by_short = {
            clean_scenario_name(col).lower(): col
            for col in initial_table.columns.get_level_values(0).unique()
        }
        for scenario_clean in scenario_names_cleaned:
            matching_scenario = full_by_short.get(scenario_clean)
            if matching_scenario is None:
                raise KeyError(f"No column match found for simplified scenario '{scenario_clean}'.")

            # Filter data for this scenario
            scenario_data = initial_table.xs(key=matching_scenario, axis=1, level=0)

            # Create a formatted table
            scenario_table = pd.DataFrame()
            scenario_table["Impact Category"] = initial_table.index
            for contribution in scenario_data.columns:
                scenario_table[contribution] = scenario_data[contribution].values

            # Calculate total impact (absolute values)
            scenario_table["Total Impact"] = scenario_table.iloc[:, 1:].abs().sum(axis=1)

            # Calculate contribution percentages
            for contribution in scenario_data.columns:
                scenario_table[f"% {contribution}"] = (
                    scenario_table[contribution] / scenario_table["Total Impact"] * 100
                )

            scenario_tables[scenario_clean] = scenario_table

        return scenario_tables
    except Exception as e:
        print(f"Error generating scenario tables: {e}")
        return None



def place_title_above_texts(fig, ax, title, texts, title_kwargs, gap_pt=8):
    """
    Place le titre juste au-dessus du plus haut texte (totaux ecrits au-dessus des barres),
    au lieu d'un decalage fixe : le titre ne chevauche jamais les totaux, quelle que soit
    la taille de police ou la longueur des nombres.
    """
    ax.set_title(title, pad=gap_pt, **title_kwargs)
    if not texts:
        return
    try:
        fig.canvas.draw()
        renderer = fig.canvas.get_renderer()
        to_axes = ax.transAxes.inverted()
        top = 1.0
        for t in texts:
            bb = t.get_window_extent(renderer)
            top = max(top, to_axes.transform((bb.x0, bb.y1))[1])
        ax.set_title(title, pad=gap_pt, y=top, **title_kwargs)
    except Exception:
        pass


def plot_comparison_bar_chart(total_impact_table):
    """
    Plots 3 figures:
    1) Main comparison bar chart
    2) Separate legend
    3) Bar chart with total values labeled
    Returns a list of matplotlib figures.
    """
    try:
        global scenario_colors, font_styles  # <== pour accéder aux styles depuis main()

        # Extraction des styles personnalisés
        fontfamily = font_styles["family"]
        fontweight = font_styles["weight"]
        fontstyle = font_styles["style"]
        title_size = font_styles["title_size"]
        label_size = font_styles["label_size"]
        legend_size = font_styles["legend_size"]
        text_size = font_styles["text_size"]

        figs = []

        # Normalize data
        normalized_data = total_impact_table.copy()
        for category in normalized_data.index:
            max_value = normalized_data.loc[category].max()
            normalized_data.loc[category] = (normalized_data.loc[category] / max_value) * 100

        categories = normalized_data.index
        scenarios = normalized_data.columns
        num_scenarios = len(scenarios)
        bar_width = 0.9 / num_scenarios
        x_positions = np.arange(len(categories))

        # Handle scenario colors
        scenario_color_map = None
        generic_colors = OKABE_ITO

        if scenario_colors is not None:
            if isinstance(scenario_colors, dict):
                scenario_color_map = scenario_colors
            elif hasattr(scenario_colors, "iterrows"):
                scenario_color_map = {
                    row["Scenario"].strip().lower(): row["Hex Code"]
                    for _, row in scenario_colors.iterrows()
                }

        # Figure 1: Main bar chart (no totals)
        fig1, ax1 = plt.subplots(figsize=(12, 7), dpi=300)
        bars = []
        labels = []

        for i, scenario in enumerate(scenarios):
            key = scenario.strip().lower()
            color = (
                scenario_color_map.get(key, palette_color(i))
                if scenario_color_map else generic_colors[i % len(generic_colors)]
            )
            bar = ax1.bar(
                x_positions + i * bar_width,
                normalized_data[scenario],
                bar_width,
                color=color,
            )
            bars.append(bar[0])
            labels.append(scenario)

        ax1.set_title("Comparison of Scenarios", fontsize=title_size, fontweight=fontweight,
              fontstyle=fontstyle, family=fontfamily, pad=30)

        ax1.set_ylabel("(%)", fontsize=label_size, fontweight=fontweight,
                       fontstyle=fontstyle, family=fontfamily)
        ax1.set_ylim(0, 100)
        ax1.set_xlim(-0.5, len(categories))
        ax1.set_xticks(x_positions + (len(scenarios) - 1) * bar_width / 2)
        ax1.set_xticklabels(categories, rotation=45, ha="right", fontsize=label_size,
                            fontweight=fontweight, fontstyle=fontstyle, family=fontfamily)
        ax1.grid(axis="y", linestyle="--", alpha=0.7)
        ax1.tick_params(axis='y', labelsize=label_size)

        plt.tight_layout()
        figs.append(fig1)
        plt.close(fig1)

        # Figure 2: Legend
        fig2, ax2 = plt.subplots(figsize=(5, 3), dpi=300)
        ax2.axis("off")
        ax2.legend(bars, labels, title="Scenarios", fontsize=legend_size, title_fontsize=legend_size, loc="center")
        plt.tight_layout()
        figs.append(fig2)
        plt.close(fig2)

        # Figure 3: With totals
        fig3, ax3 = plt.subplots(figsize=(12, 7), dpi=300)
        for i, scenario in enumerate(scenarios):
            key = scenario.strip().lower()
            color = (
                scenario_color_map.get(key, palette_color(i))
                if scenario_color_map else generic_colors[i % len(generic_colors)]
            )
            ax3.bar(
                x_positions + i * bar_width,
                normalized_data[scenario],
                bar_width,
                color=color,
            )

        total_texts = []
        for i, scenario in enumerate(scenarios):
            for j, value in enumerate(normalized_data[scenario]):
                total_value = total_impact_table.loc[categories[j], scenario]
                formatted_value = format_total(total_value, categories[j])
                total_texts.append(ax3.text(
                    x_positions[j] + i * bar_width,
                    value + 2,
                    formatted_value,
                    ha="center",
                    va="bottom",
                    fontsize=text_size,
                    color="black",
                    fontweight=fontweight,
                    fontstyle=fontstyle,
                    family=fontfamily,
                    rotation=90,
                ))

        ax3.set_ylabel("(%)", fontsize=label_size, fontweight=fontweight,
                       fontstyle=fontstyle, family=fontfamily)
        ax3.set_ylim(0, 100)
        ax3.set_xlim(-0.5, len(categories))
        ax3.set_xticks(x_positions + (len(scenarios) - 1) * bar_width / 2)
        ax3.set_xticklabels(categories, rotation=45, ha="right", fontsize=label_size,
                            fontweight=fontweight, fontstyle=fontstyle, family=fontfamily)
        ax3.grid(axis="y", linestyle="--", alpha=0.7)
        ax3.tick_params(axis='y', labelsize=label_size)

        fig3.tight_layout()
        # Titre place au-dessus des totaux (qui depassent 100 %) : plus de chevauchement
        place_title_above_texts(
            fig3, ax3, "Comparison of Scenarios (with Totals)", total_texts,
            dict(fontsize=title_size, fontweight=fontweight, fontstyle=fontstyle, family=fontfamily)
        )
        figs.append(fig3)
        plt.close(fig3)

        return figs

    except Exception as e:
        st.error(f"Error generating comparison bar chart: {str(e)}")
        return []




def plot_relative_contribution_by_scenario(scenario_tables, contributions_order):

    """
    Génère une liste de (main_figure, legend_figure) pour chaque scénario,
    affichant un graphique en barres empilées (verticales) des contributions relatives.
    """
    try:
        global contributions_colors, font_styles

        # 🔠 Extraction du style de police
        fontfamily = font_styles["family"]
        fontweight = font_styles["weight"]
        fontstyle = font_styles["style"]
        title_size = font_styles["title_size"]
        label_size = font_styles["label_size"]
        legend_size = font_styles["legend_size"]
        text_size = font_styles["text_size"]

        figures = []

        # Identifier toutes les contributions présentes
        all_contributions = set()
        for table in scenario_tables.values():
            contribution_columns = [f"% {c}" for c in contributions_order if f"% {c}" in table.columns]
            # 🔒 Skip scénario si aucune contribution n’est présente
            if not contribution_columns:
                 continue


            all_contributions.update(contribution_columns)

        # Générer la palette de couleurs
        contribution_colors = {}
        generic_colors = OKABE_ITO

        if contributions_colors is not None:
            if isinstance(contributions_colors, dict):
                for i, contrib in enumerate(sorted(all_contributions)):
                    key = contrib.replace("%", "").strip().lower()
                    contribution_colors[contrib] = contributions_colors.get(key, generic_colors[i % len(generic_colors)])
            elif hasattr(contributions_colors, "iterrows"):
                for i, contrib in enumerate(sorted(all_contributions)):
                    contrib_clean = contrib.replace("%", "").strip().lower()
                    match = contributions_colors[
                        contributions_colors["Contributions"].str.lower() == contrib_clean
                    ]
                    if not match.empty:
                        contribution_colors[contrib] = match["Hex Code"].iloc[0]
                    else:
                        contribution_colors[contrib] = generic_colors[i % len(generic_colors)]
        else:
            for i, contrib in enumerate(sorted(all_contributions)):
                contribution_colors[contrib] = generic_colors[i % len(generic_colors)]

        # Génération des graphiques
        for scenario_name, table in scenario_tables.items():
            main_fig, main_ax = plt.subplots(figsize=(16, 8), dpi=300)
            legend_fig, legend_ax = plt.subplots(figsize=(5, 3), dpi=300)
            
            categories = table["Impact Category"]
            contribution_columns = [
    f"% {c}" for c in contributions_order if f"% {c}" in table.columns
]

            percentage_data = table[contribution_columns].astype(float)
            total_impact = table["Total Impact"].astype(float)
            bar_width = 0.6
            bars = []
            labels = []
            
            bottom_pos = np.zeros(len(categories))
            bottom_neg = np.zeros(len(categories))

            for contribution in contribution_columns:
                values = percentage_data[contribution].values
                pos_values = np.where(values > 0, values, 0)
                neg_values = np.where(values < 0, values, 0)

                if np.any(pos_values):
                    bar = main_ax.bar(
                        categories, pos_values, bar_width,
                        bottom=bottom_pos,
                        color=contribution_colors[contribution]
                    )
                    bars.append(bar[0])
                    labels.append(contribution.replace("%", ""))
                    bottom_pos += pos_values

                if np.any(neg_values):
                    main_ax.bar(
                        categories, neg_values, bar_width,
                        bottom=bottom_neg,
                        color=contribution_colors[contribution]
                    )
                    bottom_neg += neg_values

            main_ax.axhline(0, color="black", linestyle="--", alpha=0.7)

            # Ajouter les valeurs totales au-dessus
            total_texts = []
            for i, total in enumerate(total_impact):
                formatted_value = format_total(total, list(categories)[i])
                total_texts.append(main_ax.text(
                    i, max(bottom_pos[i], 0) + 1,
                    formatted_value,
                    ha="center",
                    va="bottom",
                    fontsize=text_size,
                    color="black",
                    fontweight=fontweight,
                    fontstyle=fontstyle,
                    family=fontfamily,
                    rotation=90
                ))
            main_ax.set_ylabel("(%)", fontsize=label_size, fontweight=fontweight,
                               fontstyle=fontstyle, family=fontfamily)
            main_ax.set_ylim(min(bottom_neg) - 5, max(bottom_pos))
            main_ax.set_xticks(range(len(categories)))
            main_ax.set_xticklabels(
                categories,
                rotation=45,
                ha="right",
                fontsize=label_size,
                fontweight=fontweight,
                fontstyle=fontstyle,
                family=fontfamily
            )
            main_ax.grid(axis="y", linestyle="--", alpha=0.7)
            main_ax.tick_params(axis='y', labelsize=label_size)
            main_fig.tight_layout()
            # Titre au-dessus des totaux (au lieu d'un decalage fixe pad=70)
            place_title_above_texts(
                main_fig, main_ax, f"Relative Contributions - {scenario_name.capitalize()}", total_texts,
                dict(fontsize=title_size, fontweight=fontweight, fontstyle=fontstyle, family=fontfamily)
            )

            legend_ax.axis("off")
            legend_ax.legend(bars, labels, title="Contributions", fontsize=legend_size, title_fontsize=legend_size, loc="center")
            legend_fig.tight_layout()

            figures.append((main_fig, legend_fig))
            

        return figures

    except Exception as e:
        st.error(f"Error generating relative contributions: {str(e)}")
        return []





def plot_relative_contribution_by_scenario_horizontal(scenario_tables, contributions_order):

    """
    Génère des (main_figure, legend_figure) pour chaque scénario
    affichant un bar chart horizontal empilé des contributions relatives.
    """
    try:
        global contributions_colors, font_styles

        # 🔠 Styles de police
        fontfamily = font_styles["family"]
        fontweight = font_styles["weight"]
        fontstyle = font_styles["style"]
        title_size = font_styles["title_size"]
        label_size = font_styles["label_size"]
        legend_size = font_styles["legend_size"]
        text_size = font_styles["text_size"]

        figures = []

        # Identifier toutes les contributions
        all_contributions = set()
        for table in scenario_tables.values():
            
            contribution_columns = [f"% {c}" for c in contributions_order if f"% {c}" in table.columns]
            # 🔒 Skip scénario si aucune contribution n’est présente
            if not contribution_columns:
                           continue

                
            all_contributions.update(contribution_columns)

        # Générer la palette de couleurs
        contribution_color_map = {}
        generic_colors = OKABE_ITO

        if contributions_colors is not None:
            if isinstance(contributions_colors, dict):
                for i, contrib in enumerate(sorted(all_contributions)):
                    key = contrib.replace("%", "").strip().lower()
                    contribution_color_map[contrib] = contributions_colors.get(key, generic_colors[i % len(generic_colors)])
            elif hasattr(contributions_colors, "iterrows"):
                for i, contrib in enumerate(sorted(all_contributions)):
                    contrib_clean = contrib.replace("%", "").strip().lower()
                    match = contributions_colors[
                        contributions_colors["Contributions"].str.lower() == contrib_clean
                    ]
                    if not match.empty:
                        contribution_color_map[contrib] = match["Hex Code"].iloc[0]
                    else:
                        contribution_color_map[contrib] = generic_colors[i % len(generic_colors)]
        else:
            for i, contrib in enumerate(sorted(all_contributions)):
                contribution_color_map[contrib] = generic_colors[i % len(generic_colors)]

        # Génération des graphiques
        for scenario_name, table in scenario_tables.items():
            main_fig, main_ax = plt.subplots(figsize=(10, 12), dpi=300)
            legend_fig, legend_ax = plt.subplots(figsize=(5, 3), dpi=300)

            categories = table["Impact Category"]
            contribution_columns = [
    f"% {c}" for c in contributions_order if f"% {c}" in table.columns
]

            percentage_data = table[contribution_columns].astype(float)
            total_impact = table["Total Impact"].astype(float)
            bar_height = 0.6
            bars = []
            labels = []
            
            left_pos = np.zeros(len(categories))
            left_neg = np.zeros(len(categories))

            for contribution in contribution_columns:
                values = percentage_data[contribution].values
                pos_values = np.where(values > 0, values, 0)
                neg_values = np.where(values < 0, values, 0)

                if np.any(pos_values):
                    bar = main_ax.barh(
                        categories,
                        pos_values,
                        bar_height,
                        left=left_pos,
                        color=contribution_color_map[contribution]
                    )
                    bars.append(bar[0])
                    labels.append(contribution.replace("%", ""))
                    left_pos += pos_values

                if np.any(neg_values):
                    main_ax.barh(
                        categories,
                        neg_values,
                        bar_height,
                        left=left_neg,
                        color=contribution_color_map[contribution]
                    )
                    left_neg += neg_values

            main_ax.axvline(0, color="black", linestyle="--", alpha=0.7)

            # Affichage des valeurs totales à droite
            for i, total in enumerate(total_impact):
                formatted_value = format_total(total, list(categories)[i])
                main_ax.text(
                    max(left_pos[i], 0) + 5,
                    i,
                    formatted_value,
                    ha="left",
                    va="center",
                    fontsize=text_size,
                    color="black",
                    fontweight=fontweight,
                    fontstyle=fontstyle,
                    family=fontfamily
                )

            main_ax.set_title(
                f"Relative Contributions - {scenario_name.capitalize()}",
                fontsize=title_size,
                pad=30,
                fontweight=fontweight,
                fontstyle=fontstyle,
                family=fontfamily
            )
            main_ax.set_xlabel("(%)", fontsize=label_size,
                               fontweight=fontweight, fontstyle=fontstyle, family=fontfamily)
            main_ax.set_xlim(min(left_neg) - 5, max(left_pos))
            main_ax.set_yticks(range(len(categories)))
            main_ax.set_yticklabels(
                categories,
                fontsize=label_size,
                fontweight=fontweight,
                fontstyle=fontstyle,
                family=fontfamily
            )
            main_ax.grid(axis="x", linestyle="--", alpha=0.7)
            main_ax.tick_params(axis='x', labelsize=label_size)
            main_fig.tight_layout()

            legend_ax.axis("off")
            legend_ax.legend(bars, labels, title="Contributions",
                             fontsize=legend_size,title_fontsize=legend_size,  loc="center")
            legend_fig.tight_layout()

            figures.append((main_fig, legend_fig))
           

        return figures

    except Exception as e:
        st.error(f"Error generating horizontal contributions: {str(e)}")
        return []




def plot_stacked_bar_by_category(initial_table, total_impact_table, contributions_order):

    """
    Génère (main_figure, legend_figure, category_name) pour chaque catégorie d’impact.
    Utilise les valeurs réelles (initial_table). Adapté pour affichage dans Streamlit via st.pyplot().
    """
    try:
        global contributions_colors, font_styles

        # 🔠 Styles de police
        fontfamily = font_styles["family"]
        fontweight = font_styles["weight"]
        fontstyle = font_styles["style"]
        title_size = font_styles["title_size"]
        label_size = font_styles["label_size"]
        legend_size = font_styles["legend_size"]
        text_size = font_styles["text_size"]

        figures = []

        if initial_table.empty or total_impact_table.empty:
            st.warning("⚠ Données vides pour les graphiques par catégorie.")
            return []

        # 🔹 Obtenir les scénarios et contributions
        scenarios = initial_table.columns.get_level_values(0).unique()
        contributions = initial_table.columns.get_level_values(1).unique()
        categories = initial_table.index

        # 🔹 Nettoyage des noms de scénarios
        scenario_labels = {scenario: clean_scenario_name(scenario).lower() for scenario in scenarios}
        total_impact_labels = {sc.lower(): sc for sc in total_impact_table.columns}

        # 🔹 Générer le mapping de couleurs des contributions
        contrib_color_map = {}
        generic_colors = OKABE_ITO
        if contributions_colors is not None:
            if isinstance(contributions_colors, dict):
                for i, contrib in enumerate(contributions):
                    key = contrib.replace("%", "").strip().lower()
                    contrib_color_map[contrib] = contributions_colors.get(key, generic_colors[i % len(generic_colors)])
            elif hasattr(contributions_colors, "iterrows"):
                for i, contrib in enumerate(contributions):
                    contrib_clean = contrib.replace("%", "").strip().lower()
                    match = contributions_colors[contributions_colors["Contributions"].str.lower() == contrib_clean]
                    if not match.empty:
                        contrib_color_map[contrib] = match["Hex Code"].iloc[0]
                    else:
                        contrib_color_map[contrib] = generic_colors[i % len(generic_colors)]
        else:
            for i, contrib in enumerate(contributions):
                contrib_color_map[contrib] = generic_colors[i % len(generic_colors)]

        # 🔹 Générer les graphiques pour chaque catégorie
        nombre_scenario=len(scenarios)
        for category in categories:
            category_data = initial_table.loc[category]
            x_positions = np.arange(len(scenarios))*0.7
            bar_width = 0.5 # réduit l'espace entre les barres

            main_fig, main_ax = plt.subplots(figsize=(nombre_scenario*0.9, 4), dpi=300)
            legend_fig, legend_ax = plt.subplots(figsize=(5, 3), dpi=300)

            bottom_pos = np.zeros(len(scenarios))
            bottom_neg = np.zeros(len(scenarios))
            bars = []
            labels = []

            for contribution in contributions_order:
                if contribution not in contributions:
                    continue  # ignorer si non trouvé
                values = []
                for s in scenarios:
                    try:
                        val = category_data[s, contribution]
                    except KeyError:
                        val = 0
                    values.append(val)
                values = np.array(values, dtype=float)

                pos_values = np.where(values > 0, values, 0)
                neg_values = np.where(values < 0, values, 0)

                if np.any(pos_values):
                    bar = main_ax.bar(
                        x_positions,
                        pos_values,
                        width=bar_width,
                        bottom=bottom_pos,
                        color=contrib_color_map.get(contribution, "#000000")
                    )
                    bars.append(bar[0])
                    labels.append(contribution.replace("%", ""))
                    bottom_pos += pos_values

                if np.any(neg_values):
                    main_ax.bar(
                        x_positions,
                        neg_values,
                        width=bar_width,
                        bottom=bottom_neg,
                        color=contrib_color_map.get(contribution, "#000000")
                    )
                    bottom_neg += neg_values

            # ✅ Calcul de l'échelle optimisée après empilement
            all_stacked_values = list(bottom_pos) + list(bottom_neg)
            min_val = min(all_stacked_values + [0])
            max_val = max(all_stacked_values + [0])
            y_margin = 0.05 * (max_val - min_val) if max_val != min_val else 1
            ylim_min = min_val - y_margin
            ylim_max = max_val + y_margin
            main_ax.set_ylim(ylim_min, ylim_max)

            # 🔹 Ajouter les totaux
            for i, scenario in enumerate(scenarios):
                scen_clean = scenario_labels[scenario]
                if scen_clean in total_impact_labels:
                    try:
                        total_value = total_impact_table.loc[category, total_impact_labels[scen_clean]]
                        formatted = format_total(total_value, category)
                        main_ax.text(
                            x_positions[i],
                            max(bottom_pos[i], 0) + y_margin * 0.3,
                            formatted,
                            ha='center',
                            va='bottom',
                            fontsize=text_size,
                            color='black',
                            fontweight=fontweight,
                            fontstyle=fontstyle,
                            family=fontfamily
                        )
                    except KeyError:
                        continue

            # 🔹 Mise en forme
    
            main_ax.set_ylabel(f"{category}", fontsize=label_size,
                               fontweight=fontweight, fontstyle=fontstyle, family=fontfamily)
            main_ax.set_xticks(x_positions)
            main_ax.set_xticklabels(
                [scenario_labels[s] for s in scenarios],
                rotation=45,
                ha="right",
                fontsize=label_size,
                fontweight=fontweight,
                fontstyle=fontstyle,
                family=fontfamily
            )

            main_ax.set_xlim(x_positions[0] - (bar_width / 2), x_positions[-1] + (bar_width / 2))
            main_ax.grid(axis="y", linestyle="--", alpha=0.7)
            main_ax.tick_params(axis='y', labelsize=label_size)
            main_fig.tight_layout()

            legend_ax.axis("off")
            legend_ax.legend(bars, labels, title="Contributions",
                             fontsize=legend_size, title_fontsize=legend_size, loc="center")
            legend_fig.tight_layout()

            figures.append((main_fig, legend_fig, category))

        return figures

    except Exception as e:
        st.error(f"❌ Error generating graph per category : {str(e)}")
        return []






def finalize_fixed_axes(fig, ax, ax_w_in, ax_h_in, pad_in=0.15):
    """
    Fixe la taille physique du diagramme (l'axes) a ax_w_in x ax_h_in pouces,
    puis agrandit la figure pour englober tout le contenu (labels, totaux)
    sans jamais retrecir le diagramme. La taille du diagramme reste donc
    identique entre 'Main Chart' et 'Chart with Totals' ; seule l'image
    entiere grandit (surtout en hauteur) selon la taille du texte.
    """
    try:
        # Figure temporaire genereuse pour ne rien clipper pendant la mesure
        tw, th = ax_w_in + 10.0, ax_h_in + 10.0
        fig.set_size_inches(tw, th)
        ax.set_position([5.0 / tw, 5.0 / th, ax_w_in / tw, ax_h_in / th])
        fig.canvas.draw()
        renderer = fig.canvas.get_renderer()
        # Bounding box de tout le contenu (en pouces)
        tb = fig.get_tightbbox(renderer)
        ab = ax.get_window_extent(renderer).transformed(fig.dpi_scale_trans.inverted())
        new_w = tb.width + 2 * pad_in
        new_h = tb.height + 2 * pad_in
        left_in = (ab.x0 - tb.x0) + pad_in
        bottom_in = (ab.y0 - tb.y0) + pad_in
        # La police est en points (taille absolue) : agrandir la figure ne
        # change pas le texte, donc le diagramme garde exactement ax_w_in x ax_h_in
        fig.set_size_inches(new_w, new_h)
        ax.set_position([left_in / new_w, bottom_in / new_h, ax_w_in / new_w, ax_h_in / new_h])
    except Exception:
        fig.tight_layout()
    return fig


def default_combined_bar_layout(num_categories, num_scenarios, orientation="vertical"):
    """
    Valeurs par defaut (en cm) de l'epaisseur des barres et de l'espace entre
    categories d'impact, calculees pour reproduire l'ancien rendu.
    """
    num_categories = max(1, num_categories)
    num_scenarios = max(1, num_scenarios)
    if orientation == "vertical":
        # Ancien rendu : diagramme de 14 in, 1.5 unite par categorie, barres de 0.9/n
        unit_cm = 14.0 * 2.54 / (num_categories * 1.5)
        bar_cm = (0.9 / num_scenarios) * unit_cm
        gap_cm = 0.6 * unit_cm
    else:
        # Ancien rendu horizontal : 2.0 unites par categorie
        bar_h = min(0.9, max(0.3, 1.5 / num_scenarios))
        fig_height = max(6, num_categories * 2.0 * 0.7)
        ax_h_in = max(4.0, fig_height - 1.2)
        unit_cm = ax_h_in * 2.54 / (num_categories * 2.0)
        bar_cm = bar_h * unit_cm
        gap_cm = max(0.1, 2.0 - num_scenarios * bar_h) * unit_cm
    return round(bar_cm, 2), round(gap_cm, 2)


def place_total_labels(fig, ax, labels, orientation, axis_len_in, mode="smart", offset_pt=6.0):
    """
    Place les totaux SANS modifier l'axe (qui reste borne a 100 %).
    - mode "smart" : chaque total est colle a sa barre (decalage fixe en points) ;
      s'il devait chevaucher la bordure du cadre, il est remonte juste au-dessus
      du cadre (ou a droite en horizontal) pour ne jamais etre barre.
    - mode "row" : tous les totaux sont alignes sur une ligne juste au-dessus
      du cadre (a droite en horizontal).
    """
    if not labels:
        return
    renderer = fig.canvas.get_renderer()
    lo, hi = ax.get_ylim() if orientation == "vertical" else ax.get_xlim()
    pt_per_unit = axis_len_in * 72.0 / (hi - lo)
    for ann in labels:
        # Ecart (en points) entre la barre - ou la bordure du cadre - et le texte
        ann.xyann = (0, offset_pt) if orientation == "vertical" else (offset_pt, 0)
        x, y = ann.xy
        v = y if orientation == "vertical" else x
        bb = ann.get_window_extent(renderer)
        ext_pt = (bb.height if orientation == "vertical" else bb.width) * 72.0 / fig.dpi
        # Fin du texte (en unites de donnees) s'il reste colle a la barre
        end_v = v + (offset_pt + ext_pt) / pt_per_unit
        if mode == "row" or end_v > hi:
            v = hi  # demarre juste apres la bordure du cadre
        ann.xy = (x, v) if orientation == "vertical" else (v, y)


def default_combined_bar_length_cm(num_scenarios, orientation="vertical"):
    """
    Longueur par defaut (en cm) de l'echelle des barres (axe 0-100 %) :
    hauteur du diagramme en vertical, largeur en horizontal (ancien rendu).
    """
    if orientation == "vertical":
        return round(8.0 * 2.54, 1)
    return round(max(6.0, 9.0 + 0.5 * max(1, num_scenarios)) * 2.54, 1)


def plot_combined_graph_with_scenario_hatches(percentage_table, total_impact_table, contributions_order,
                                              bar_width_cm=None, category_gap_cm=None,
                                              bar_length_cm=None, show_total_pct=True,
                                              totals_mode="smart", totals_offset_pt=6.0):


    """
    Génère des graphiques combinés par scénario avec hachures et légendes.
    Retourne une liste de tuples : (description, figure).
    """
    try:
        global contributions_colors, scenario_colors, font_styles
        figures = []

        # Récupération des styles de police
        fontfamily = font_styles["family"]
        fontweight = font_styles["weight"]
        fontstyle = font_styles["style"]
        title_size = font_styles["title_size"]
        label_size = font_styles["label_size"]
        legend_size = font_styles["legend_size"]
        text_size = font_styles["text_size"]

        if percentage_table.empty or total_impact_table.empty:
            st.warning("Missing data for the combined scenario chart.")
            return []

        categories = percentage_table.index
        scenarios = percentage_table.columns.get_level_values(0).unique()
        contributions = percentage_table.columns.get_level_values(1).unique()
        num_scenarios = len(scenarios)
        # Epaisseur des barres et espace entre categories (en cm, 1 unite = 1 cm)
        def_bar, def_gap = default_combined_bar_layout(len(categories), num_scenarios, "vertical")
        bar_width = float(bar_width_cm) if bar_width_cm else def_bar
        category_gap = float(category_gap_cm) if category_gap_cm is not None else def_gap
        group_width = num_scenarios * bar_width
        x_positions = np.arange(len(categories)) * (group_width + category_gap)
        # Limites de l'axe : demi-espace de chaque cote
        x_min = -bar_width / 2 - category_gap / 2
        x_max = x_positions[-1] + group_width - bar_width / 2 + category_gap / 2
        ax_width_in = max(2.0, (x_max - x_min) / 2.54)

        # Définir les hachures par scénario
        scenario_hatch_dict = {}
        for scenario in scenarios:
              key = clean_scenario_name(scenario).lower()
              scenario_hatch_dict[scenario] = scenario_hatches.get(key, None)


        # Gérer les couleurs des contributions
        contrib_color_map = {}
        if contributions_colors is not None:
            if isinstance(contributions_colors, dict):
                for i, contrib in enumerate(contributions):
                    key = contrib.replace("%", "").strip().lower()
                    contrib_color_map[contrib] = contributions_colors.get(key, palette_color(i))
            elif hasattr(contributions_colors, "iterrows"):
                for i, contrib in enumerate(contributions):
                    key = contrib.replace("%", "").strip().lower()
                    match = contributions_colors[
                        contributions_colors["Contributions"].str.lower() == key
                    ]
                    contrib_color_map[contrib] = match["Hex Code"].iloc[0] if not match.empty else palette_color(i)
        else:
            for i, contrib in enumerate(contributions):
                contrib_color_map[contrib] = palette_color(i)

        # Gérer les couleurs des scénarios
        scenario_color_map = {}
        if scenario_colors is not None:
            if isinstance(scenario_colors, dict):
                for scen in scenarios:
                    key = clean_scenario_name(scen).lower()
                    scenario_color_map[scen] = scenario_colors.get(key, "#FFFFFF")
            elif hasattr(scenario_colors, "iterrows"):
                for scen in scenarios:
                    key = clean_scenario_name(scen).lower()
                    match = scenario_colors[
                        scenario_colors["Scenario"].str.lower() == key
                    ]
                    scenario_color_map[scen] = match["Hex Code"].iloc[0] if not match.empty else "#FFFFFF"
        else:
            for scen in scenarios:
                scenario_color_map[scen] = "#FFFFFF"

        # Fonction interne pour créer un graphique
        def create_figure(show_totals=False):
            fig, ax = plt.subplots(figsize=(16, 10), dpi=300)
            total_labels = []
            ax.tick_params(axis='y', labelsize=label_size)

            for i, scenario in enumerate(scenarios):
                bottom_pos = np.zeros(len(categories))
                bottom_neg = np.zeros(len(categories))

                for contrib in [c for c in contributions_order if c in contributions]:

                    try:
                        values = (
                            percentage_table.xs((scenario, contrib), axis=1, level=[0, 1])
                            .values.flatten()
                        )
                    except KeyError:
                        continue

                    values = np.nan_to_num(values.astype(float))
                    pos_values = np.where(values > 0, values, 0)
                    neg_values = np.where(values < 0, values, 0)

                    if np.any(pos_values):
                        ax.bar(
                            x_positions + i * bar_width,
                            pos_values,
                            bar_width,
                            bottom=bottom_pos,
                            color=contrib_color_map[contrib],
                            hatch=scenario_hatch_dict.get(scenario, None)
,
                            edgecolor='black'
                        )
                        bottom_pos += pos_values

                    if np.any(neg_values):
                        ax.bar(
                            x_positions + i * bar_width,
                            neg_values,
                            bar_width,
                            bottom=bottom_neg,
                            color=contrib_color_map[contrib],
                            hatch=scenario_hatch_dict.get(scenario, None)
,
                            edgecolor='black'
                        )
                        bottom_neg += neg_values

                if show_totals:
                    scenario_clean = clean_scenario_name(scenario)
                    try:
                        totals = total_impact_table[scenario_clean].values
                        cat_max_values = total_impact_table.max(axis=1).values
                        for j, total in enumerate(totals):
                            formatted = format_total(total, total_impact_table.index[j])
                            max_val = cat_max_values[j]
                            pct = (total / max_val * 100) if max_val != 0 else 0
                            label = f"{formatted} ({pct:.0f}%)" if show_total_pct else formatted

                            total_labels.append(ax.annotate(
                                label,
                                xy=(x_positions[j] + i * bar_width, bottom_pos[j]),
                                xytext=(0, 3), textcoords='offset points',
                                ha='center',
                                va='bottom',
                                rotation=90,
                                fontsize=text_size,
                                fontweight=fontweight,
                                fontstyle=fontstyle,
                                family=fontfamily,
                                annotation_clip=False
                            ))
                    except KeyError:
                        pass

            ax.set_xticks(x_positions + (num_scenarios - 1) * bar_width / 2)
            ax.set_xticklabels(categories, rotation=45, ha='right',
                               fontsize=label_size, fontweight=fontweight,
                               fontstyle=fontstyle, family=fontfamily)
            ax.set_ylabel("Contribution (%)", fontsize=label_size,
                          fontweight=fontweight, fontstyle=fontstyle, family=fontfamily)
            ax.axhline(0, color='black', linestyle='--', alpha=0.5)
            ax.set_ylim(top=100)
            ax.set_xlim(x_min, x_max)
            ax_height_in = max(1.0, (float(bar_length_cm) if bar_length_cm else default_combined_bar_length_cm(num_scenarios, "vertical")) / 2.54)
            place_total_labels(fig, ax, total_labels, "vertical", ax_height_in, mode=totals_mode, offset_pt=totals_offset_pt)
            finalize_fixed_axes(fig, ax, ax_width_in, ax_height_in)
            return fig

        # Graphiques principaux
        figures.append(("Main Chart", create_figure(show_totals=False)))
        figures.append(("Chart with Totals", create_figure(show_totals=True)))


        # Légende pour les contributions
        legend_fig1, ax1 = plt.subplots(figsize=(6, 2), dpi=300)
        ax1.axis('off')

        contrib_handles = [
    plt.Rectangle((0, 0), 1, 1, facecolor=color, edgecolor='black')
    for _, color in contrib_color_map.items()
]
        contrib_labels = list(contrib_color_map.keys())

        legend1 =ax1.legend(
    contrib_handles,
    contrib_labels,
    title="Contributions",
    fontsize=legend_size,
    title_fontsize=legend_size,
    ncol=2,
    loc='center'
)
        legend1.get_frame().set_linewidth(0) 
        
        figures.append(("Contributions Legend", legend_fig1))


        # Légende pour les scénarios
        legend_fig2, ax2 = plt.subplots(figsize=(6, 2), dpi=300)
        ax2.axis('off')

        scenario_handles = [
        plt.Rectangle((0, 0), 1, 1,
                  facecolor='white',
                  hatch=legend_hatch(scenario_hatch_dict[scenario]),
                  edgecolor='black')
        for scenario in scenarios
]
        scenario_labels = [clean_scenario_name(s) for s in scenarios]

        legend2 = ax2.legend(
    scenario_handles,
    scenario_labels,
    title="Scenarios",
    fontsize=legend_size,
    title_fontsize=legend_size,
    ncol=2,
    loc='center'
)
        legend2.get_frame().set_linewidth(0)
        figures.append(("Scenarios Legend", legend_fig2))
        return figures

    except Exception as e:
        st.error(f"Error generating combined scenario chart: {str(e)}")
        return []
def plot_combined_graph_with_scenario_hatches_horizontal(percentage_table, total_impact_table, contributions_order,
                                                         bar_width_cm=None, category_gap_cm=None,
                                              bar_length_cm=None, show_total_pct=True,
                                              totals_mode="smart", totals_offset_pt=6.0):
    """
    Génère des graphiques horizontaux combinés par scénario avec hachures et légendes.
    Retourne une liste de tuples : (description, figure).
    """
    try:
        global contributions_colors, scenario_colors, font_styles, scenario_hatches
        figures = []

        fontfamily = font_styles["family"]
        fontweight = font_styles["weight"]
        fontstyle = font_styles["style"]
        title_size = font_styles["title_size"]
        label_size = font_styles["label_size"]
        legend_size = font_styles["legend_size"]
        text_size = font_styles["text_size"]

        if percentage_table.empty or total_impact_table.empty:
            st.warning("Missing data for the combined scenario chart.")
            return []

        categories = percentage_table.index
        scenarios = percentage_table.columns.get_level_values(0).unique()
        contributions = percentage_table.columns.get_level_values(1).unique()
        num_scenarios = len(scenarios)
        num_categories = len(categories)
        # Epaisseur des barres et espace entre categories (en cm, 1 unite = 1 cm)
        def_bar, def_gap = default_combined_bar_layout(num_categories, num_scenarios, "horizontal")
        bar_height = float(bar_width_cm) if bar_width_cm else def_bar
        category_gap = float(category_gap_cm) if category_gap_cm is not None else def_gap
        group_height = num_scenarios * bar_height
        y_positions = np.arange(num_categories) * (group_height + category_gap)
        y_min = -bar_height / 2 - category_gap / 2
        y_max = y_positions[-1] + group_height - bar_height / 2 + category_gap / 2
        ax_height_in = max(1.0, (y_max - y_min) / 2.54)
        fig_width = 12 + 0.5 * num_scenarios
        fig_height = ax_height_in + 1.2

        
        

       

        scenario_hatch_dict = {scenario: scenario_hatches.get(clean_scenario_name(scenario).lower(), None) for scenario in scenarios}

        # Couleurs contributions
        contrib_color_map = {contrib: contributions_colors.get(contrib.replace("%", "").strip().lower(), palette_color(i)) for i, contrib in enumerate(contributions)}

        # Couleurs scénarios
        scenario_color_map = {scen: scenario_colors.get(clean_scenario_name(scen).lower(), "#FFFFFF") for scen in scenarios}

        def create_figure(show_totals=False):
            fig, ax = plt.subplots(figsize=(fig_width, fig_height), dpi=300)
            total_labels = []
            ax.tick_params(axis='x', labelsize=label_size)
            x_lo, x_hi = 0.0, 0.0  # etendue sur TOUS les scenarios

            for i, scenario in enumerate(scenarios):
                left_pos = np.zeros(len(categories))
                left_neg = np.zeros(len(categories))

                for contrib in [c for c in contributions_order if c in contributions]:
                    try:
                        values = percentage_table.xs((scenario, contrib), axis=1, level=[0, 1]).values.flatten()
                    except KeyError:
                        continue
                    values = np.nan_to_num(values.astype(float))
                    pos_values = np.where(values > 0, values, 0)
                    neg_values = np.where(values < 0, values, 0)

                    if np.any(pos_values):
                        ax.barh(y_positions + i * bar_height, pos_values, bar_height, left=left_pos, color=contrib_color_map[contrib], hatch=scenario_hatch_dict.get(scenario, None), edgecolor='black')
                        left_pos += pos_values

                    if np.any(neg_values):
                        ax.barh(y_positions + i * bar_height, neg_values, bar_height, left=left_neg, color=contrib_color_map[contrib], hatch=scenario_hatch_dict.get(scenario, None), edgecolor='black')
                        left_neg += neg_values

                x_lo = min(x_lo, float(left_neg.min()))
                x_hi = max(x_hi, float(left_pos.max()))

                if show_totals:
                    scenario_clean = clean_scenario_name(scenario)
                    try:
                        totals = total_impact_table[scenario_clean].values
                        cat_max_values = total_impact_table.max(axis=1).values
                        for j, total in enumerate(totals):
                            formatted = format_total(total, total_impact_table.index[j])
                            max_val = cat_max_values[j]
                            pct = (total / max_val * 100) if max_val != 0 else 0
                            label = f"{formatted} ({pct:.0f}%)" if show_total_pct else formatted

                            total_labels.append(ax.annotate(label, xy=(left_pos[j], y_positions[j] + i * bar_height), xytext=(3, 0), textcoords='offset points', va='center', ha='left', fontsize=text_size, fontweight=fontweight, fontstyle=fontstyle, family=fontfamily, annotation_clip=False))
                    except KeyError:
                        pass

            ax.set_yticks(y_positions + (num_scenarios - 1) * bar_height / 2)
            ax.set_yticklabels(categories, fontsize=label_size, fontweight=fontweight, fontstyle=fontstyle, family=fontfamily)
            ax.set_xlabel("Contribution (%)", fontsize=label_size, fontweight=fontweight, fontstyle=fontstyle, family=fontfamily)
            ax.axvline(0, color='black', linestyle='--', alpha=0.5)
            ax.set_xlim(left=x_lo - 5, right=x_hi + 5)
            ax.set_ylim(y_min, y_max)
            ax_width_in = max(1.0, (float(bar_length_cm) if bar_length_cm else default_combined_bar_length_cm(num_scenarios, "horizontal")) / 2.54)
            place_total_labels(fig, ax, total_labels, "horizontal", ax_width_in, mode=totals_mode, offset_pt=totals_offset_pt)
            finalize_fixed_axes(fig, ax, ax_width_in, ax_height_in)
            return fig

        figures.append(("Main Chart (Horizontal)", create_figure(False)))
        figures.append(("Chart with Totals (Horizontal)", create_figure(True)))

        legend_fig1, ax1 = plt.subplots(figsize=(6, 2), dpi=300)
        ax1.axis('off')
        contrib_handles = [plt.Rectangle((0, 0), 1, 1, facecolor=color, edgecolor='black') for _, color in contrib_color_map.items()]
        contrib_labels = list(contrib_color_map.keys())
        legend1 = ax1.legend(contrib_handles, contrib_labels, title="Contributions", fontsize=legend_size, title_fontsize=legend_size, ncol=2, loc='center')
        legend1.get_frame().set_linewidth(0)
        figures.append(("Contributions Legend (Horizontal)", legend_fig1))

        legend_fig2, ax2 = plt.subplots(figsize=(6, 2), dpi=300)
        ax2.axis('off')
        scenario_handles = [plt.Rectangle((0, 0), 1, 1, facecolor='white', hatch=legend_hatch(scenario_hatch_dict[scenario]), edgecolor='black') for scenario in scenarios]
        scenario_labels = [clean_scenario_name(s) for s in scenarios]
        legend2 = ax2.legend(scenario_handles, scenario_labels, title="Scenarios", fontsize=legend_size, title_fontsize=legend_size, ncol=2, loc='center')
        legend2.get_frame().set_linewidth(0)
        figures.append(("Scenarios Legend (Horizontal)", legend_fig2))

        return figures

    except Exception as e:
        st.error(f"Error generating horizontal combined scenario chart: {str(e)}")
        return []

    




def generate_table_graph(df, table_name):
    """
    Generates a matplotlib figure displaying a DataFrame as a nicely formatted table:
    - Numeric values in scientific notation (2E)
    - First row (header) and first column highlighted with a lightblue background and bold text.
    """
    # Copy the DataFrame and format numeric values in 2E notation
    formatted_df = df.copy()
    formatted_df = formatted_df.applymap(
        lambda x: f"{x:.2E}" if isinstance(x, (int, float)) else x
    )

    num_rows, num_cols = formatted_df.shape
    fig_width = max(10, num_cols * 1.5)
    fig_height = max(6, num_rows * 0.5 + 1.5)
    fig, ax = plt.subplots(figsize=(fig_width, fig_height))
    ax.set_axis_off()

    title_y_position = 0.96
    table_y_position = 0.96 - (num_rows * 0.01)
    table_y_position = max(0.70, table_y_position)

    plt.figtext(
        0.5,
        title_y_position,
        f"Table: {table_name}",
        fontsize=18,
        fontweight="bold",
        ha="center",
    )

    table = ax.table(
        cellText=formatted_df.values,
        colLabels=formatted_df.columns,
        rowLabels=formatted_df.index,
        loc="center",
        cellLoc="center",
        colWidths=[0.5] * formatted_df.shape[1],
    )

    for key, cell in table.get_celld().items():
        cell.set_fontsize(14)
        cell.set_height(0.08)

    header_color = "lightblue"
    first_col_color = "lightblue"

    # Header row
    for i in range(len(formatted_df.columns)):
        cell = table[0, i]
        cell.set_text_props(weight="bold")
        cell.set_facecolor(header_color)

    # First column
    for i in range(len(formatted_df.index)):
        cell = table[i + 1, -1]
        cell.set_text_props(weight="bold")
        cell.set_facecolor(first_col_color)

    plt.subplots_adjust(top=table_y_position, bottom=0.1)
    return fig


def find_file_path(file_name, start_directory=None):
    """
    Recursively searches for a file starting from 'start_directory'.
    If 'start_directory' is not provided, uses the user's home directory.
    """
    if start_directory is None:
        start_directory = os.path.expanduser("~")

    for root, dirs, files in os.walk(start_directory):
        if file_name in files:
            return os.path.join(root, file_name)

    return None





###############################
#   Rendering (cache, exports, settings)
###############################

CHART_CHECKBOXES = [
    ("gen_comparison", "Generate Scenario Comparison Charts"),
    ("gen_details_v", "Generate Detailed Scenario Charts (Vertical)"),
    ("gen_details_h", "Generate Detailed Scenario Charts (Horizontal)"),
    ("gen_category", "Generate Category-by-Category Charts"),
]
COMBINED_MODES = ["Do not generate", "Vertical", "Horizontal"]
TOTALS_MODE_LABELS = {
    "smart": "Next to each bar (moved outside the frame only if it would cross it)",
    "row": "All aligned just outside the frame",
}
_HEX_RE = re.compile(r"^#[0-9a-fA-F]{6}$")
_COMBINED_DYNAMIC_KEYS = ("combined_bar_cm_", "combined_gap_cm_", "combined_len_cm_")


def safe_filename(text):
    """Nom de fichier lisible : 'GW (kg CO$_{2}$ eq)' -> 'gw_kg_co2_eq'."""
    t = re.sub(r"\$[_^]\{(.*?)\}\$", r"\1", str(text))   # indices / exposants matplotlib
    t = t.replace("\\", "").replace("$", "")
    t = re.sub(r"[^\w\-]+", "_", t).strip("_").lower()
    return t or "chart"


def figure_to_outputs(fig, name):
    """Convertit une figure en PNG 300 DPI + SVG + apercu leger, puis la ferme (memoire)."""
    out = {"name": name}
    for key, fmt, dpi in (("png", "png", EXPORT_DPI), ("svg", "svg", None), ("preview", "png", PREVIEW_DPI)):
        buf = io.BytesIO()
        if dpi:
            fig.savefig(buf, format=fmt, dpi=dpi, bbox_inches="tight")
        else:
            fig.savefig(buf, format=fmt, bbox_inches="tight")
        out[key] = buf.getvalue()
    plt.close(fig)
    return out


def _group(title, figs_and_names, ratios=None):
    return {"title": title, "ratios": ratios,
            "items": [figure_to_outputs(fig, name) for fig, name in figs_and_names]}


@st.cache_data(show_spinner=False, max_entries=10, ttl=3600)
def load_tables(data_sig, sheet_name, _file_bytes):
    """Lecture + validation du fichier, mise en cache par empreinte du fichier et feuille."""
    return analyze_excel_and_generate_tables(io.BytesIO(_file_bytes), sheet_name)


def apply_render_settings(style, total_impact_table):
    """Applique les reglages aux variables globales utilisees par les fonctions de trace."""
    global scenario_colors, contributions_colors, scenario_hatches, font_styles
    font_styles = dict(style["font"])
    scenario_colors = dict(style["scenario_colors"])
    contributions_colors = dict(style["contribution_colors"])
    scenario_hatches = {k: (None if v in (None, "", "None") else v)
                        for k, v in style["scenario_hatches"].items()}
    TOTAL_GROUPS.clear()
    TOTAL_GROUPS.update({cat: list(total_impact_table.loc[cat].values) for cat in total_impact_table.index})
    # Indices / exposants (CO2, m3...) dans la meme police que le reste du texte
    family = font_styles["family"]
    plt.rcParams["mathtext.fontset"] = "custom"
    plt.rcParams["mathtext.rm"] = family
    plt.rcParams["mathtext.it"] = f"{family}:italic"
    plt.rcParams["mathtext.bf"] = f"{family}:bold"


def _build_groups(kind, initial_table, total_impact_table, scenario_names, order, extra):
    groups = []
    if kind == "comparison":
        figs = plot_comparison_bar_chart(total_impact_table)
        meta = [("Main Scenario Comparison", "scenario_comparison"),
                ("Legend", "scenario_comparison_legend"),
                ("Scenario Comparison (with Totals)", "scenario_comparison_with_totals")]
        for fig, (title, fname) in zip(figs, meta):
            groups.append(_group(title, [(fig, fname)]))

    elif kind in ("details_v", "details_h"):
        vertical = kind == "details_v"
        scenario_tables = generate_tables_by_scenario(initial_table, scenario_names) or {}
        plot_fn = plot_relative_contribution_by_scenario if vertical else plot_relative_contribution_by_scenario_horizontal
        suffix = "vertical" if vertical else "horizontal"
        figs = plot_fn(scenario_tables, order) if scenario_tables else []
        for name, (main_fig, legend_fig) in zip(scenario_tables.keys(), figs):
            title = f"Scenario: {name}" + ("" if vertical else " (Horizontal)")
            groups.append(_group(title, [(main_fig, f"{safe_filename(name)}_{suffix}"),
                                         (legend_fig, f"{safe_filename(name)}_legend_{suffix}")],
                                 ratios=[3, 1]))

    elif kind == "category":
        figs = plot_stacked_bar_by_category(initial_table, total_impact_table, order)
        for main_fig, legend_fig, category in figs:
            name = safe_filename(category)
            groups.append(_group(f"Category: {category}",
                                 [(main_fig, f"category_{name}"), (legend_fig, f"category_{name}_legend")],
                                 ratios=[4, 1]))

    elif kind in ("combined_v", "combined_h"):
        percentage_table = generate_percentage_table(initial_table, total_impact_table)
        plot_fn = (plot_combined_graph_with_scenario_hatches if kind == "combined_v"
                   else plot_combined_graph_with_scenario_hatches_horizontal)
        figs = plot_fn(percentage_table, total_impact_table, order,
                       extra.get("bar_cm"), extra.get("gap_cm"), extra.get("length_cm"),
                       extra.get("show_pct", True), extra.get("totals_mode", "smart"),
                       extra.get("totals_offset_pt", 6.0))
        legends = [(n, f) for n, f in figs if "Legend" in n]
        charts = [(n, f) for n, f in figs if "Legend" not in n]
        for name, fig in legends + charts:     # legendes d'abord, comme avant
            groups.append(_group(name, [(fig, safe_filename(name))]))
    return groups


@st.cache_data(show_spinner=False, max_entries=40, ttl=3600)
def render_charts(kind, data_sig, sheet_name, style_json, extra_json, _file_bytes):
    """
    Genere les graphes d'un type donne et renvoie les images (PNG, SVG, apercu).
    Mis en cache : tant que le fichier et les reglages utiles a CE type de graphe ne
    changent pas, rien n'est recalcule (un clic ailleurs dans l'interface est instantane).
    """
    initial_table, _combined, total_impact_table, scenario_names = load_tables(data_sig, sheet_name, _file_bytes)
    style = json.loads(style_json)
    extra = json.loads(extra_json) if extra_json else {}
    with _RENDER_LOCK:
        apply_render_settings(style, total_impact_table)
        try:
            return _build_groups(kind, initial_table, total_impact_table, scenario_names,
                                 style["contribution_order"], extra)
        finally:
            plt.close("all")   # aucune figure ne reste en memoire


def show_group(group, key_prefix):
    """Affiche un groupe de figures (ex. graphe + legende) avec boutons PNG / SVG."""
    st.subheader(group["title"])
    items = group["items"]
    ratios = group.get("ratios")
    cols = st.columns(ratios) if ratios and len(ratios) == len(items) else [st.container() for _ in items]
    for j, (col, item) in enumerate(zip(cols, items)):
        with col:
            st.image(item["preview"])
            c1, c2 = st.columns(2)
            c1.download_button(f"📥 PNG ({EXPORT_DPI} DPI)", data=item["png"],
                               file_name=f"{item['name']}.png", mime="image/png",
                               key=f"{key_prefix}_{j}_png")
            c2.download_button("📐 SVG", data=item["svg"],
                               file_name=f"{item['name']}.svg", mime="image/svg+xml",
                               key=f"{key_prefix}_{j}_svg")
    st.markdown("---")


def build_zip(all_sections, settings):
    """ZIP de tous les graphes affiches (PNG + SVG), un dossier par type, + les reglages."""
    buf = io.BytesIO()
    used = set()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("lca_graph_settings.json", json.dumps(settings, indent=2, ensure_ascii=False))
        for header, groups in all_sections:
            folder = safe_filename(header)
            for group in groups:
                for item in group["items"]:
                    base = f"{folder}/{item['name']}"
                    path, k = base, 2
                    while path in used:
                        path, k = f"{base}_{k}", k + 1
                    used.add(path)
                    zf.writestr(path + ".png", item["png"])
                    zf.writestr(path + ".svg", item["svg"])
    return buf.getvalue()


def _num(value, lo, hi, default):
    """Nombre borne entre lo et hi, ou la valeur par defaut si invalide."""
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not np.isfinite(value):
        return float(default)
    return float(min(hi, max(lo, value)))


def _apply_uploaded_settings():
    """Callback : applique un fichier de reglages JSON aux widgets (avant leur creation)."""
    ss = st.session_state
    uploaded = ss.get("settings_upload")
    if uploaded is None:
        return
    try:
        s = json.loads(uploaded.getvalue().decode("utf-8"))
        if not isinstance(s, dict):
            raise ValueError("the file does not contain a JSON object")
    except Exception as e:
        ss["_settings_msg"] = ("error", f"Could not read the settings file: {e}")
        return

    font = s.get("font") or {}
    for key, options in (("family", FONT_FAMILIES), ("weight", FONT_WEIGHTS), ("style", FONT_STYLES)):
        if font.get(key) in options:
            ss[f"font_{key}"] = font[key]
    for name, _label, lo, hi, default in FONT_SIZE_SLIDERS:
        if name in font:
            ss[f"font_{name}"] = int(round(_num(font.get(name), lo, hi, default)))
    label_by_mode = {mode: label for label, mode in PRECISION_MODES.items()}
    if font.get("precision_mode") in label_by_mode:
        ss["precision_mode_label"] = label_by_mode[font["precision_mode"]]

    n_colors = 0
    for prefix, colors in (("scen_color::", s.get("scenario_colors") or {}),
                           ("contrib_color::", s.get("contribution_colors") or {})):
        for name, hex_code in colors.items():
            if isinstance(hex_code, str) and _HEX_RE.match(hex_code):
                ss[f"{prefix}{str(name).lower()}"] = hex_code.lower()
                n_colors += 1
    if n_colors:
        ss["customize_colors"] = True

    n_hatches = 0
    for name, hatch in (s.get("scenario_hatches") or {}).items():
        hatch = "None" if hatch in (None, "", "None") else hatch
        if hatch in HATCH_CHOICES:
            ss[f"hatch_selector_{str(name).lower()}"] = hatch
            n_hatches += 1
    if n_hatches:
        ss["customize_hatches"] = True

    if isinstance(s.get("contribution_order"), list):
        ss["_loaded_order"] = [str(c) for c in s["contribution_order"]]

    charts = s.get("charts") or {}
    for key, _label in CHART_CHECKBOXES:
        if isinstance(charts.get(key), bool):
            ss[key] = charts[key]
    if charts.get("combined_chart_mode") in COMBINED_MODES:
        ss["combined_chart_mode"] = charts["combined_chart_mode"]

    combined = s.get("combined") or {}
    if isinstance(combined.get("show_pct"), bool):
        ss["combined_show_total_pct"] = combined["show_pct"]
    if combined.get("totals_mode") in TOTALS_MODE_LABELS:
        ss["combined_totals_mode"] = TOTALS_MODE_LABELS[combined["totals_mode"]]
    if "totals_offset_pt" in combined:
        ss["combined_totals_offset"] = _num(combined.get("totals_offset_pt"), 0.0, 30.0, 6.0)
    ss["_loaded_combined"] = {k: combined.get(k) for k in ("orientation", "bar_cm", "gap_cm", "length_cm")}
    # Les cles des dimensions dependent des donnees : on les reinitialise depuis le fichier
    for key in [k for k in ss.keys() if str(k).startswith(_COMBINED_DYNAMIC_KEYS)]:
        del ss[key]

    ss["_settings_msg"] = ("success", "Settings loaded. Names that do not exist in your data are ignored.")


###############################
#           Main App
###############################

def show_introduction_page():
    st.markdown(
    """
    <h1 style='text-align: center; font-size: 3em;'>
        LCA Graph Dashboard 🌍
    </h1>
    <p style='text-align: center; font-size: 1.1em;'>

    """,
    unsafe_allow_html=True
)

    st.markdown("""
    This tool is designed to automatically generate insightful **Life Cycle Assessment (LCA)** graphs from a structured Excel file.

    It combines a simple data entry template with a Python-powered interface to visualize your environmental impact data across multiple dimensions, including:
    - 📊 **Scenario comparison**
    - 🧩 **Relative contribution charts (vertical & horizontal)**
    - 📚 **Impact analysis by category**
    - 🎯 **Combined views of scenarios and contributions**
    """)

    st.markdown("---")

    # ----------------------------------------
    # 🔹 Étape 1 : Télécharger la template
    # ----------------------------------------
    st.header("Step 1: Prepare your data")

    st.markdown("""
    Download the Excel template below and fill in your LCA data accordingly.

    **Structuring rules:**
    - You can define multiple scenarios.
    - Each scenario can include multiple contributions.
    - A contribution can appear in several scenarios.
    - You can include as many impact categories as needed.
    - Each scenario needs a **different short name** (the text in brackets, e.g. *Scenario 1 (**baseline**)*).

    🔎 **Note**: The Excel file already contains examples data to help you visualize the kind of charts that will be generated.
    """)

    with open("assets/LCA_template.xlsx", "rb") as f:
        excel_bytes = f.read()

    st.download_button(
        label="📥 Download Excel Template",
        data=excel_bytes,
        file_name="LCA_template.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    st.markdown("---")


def show_generator_page():
    ss = st.session_state

    st.markdown("## 📊 Step 2: Upload and Generate Graphs")

    # Reglages : charger en haut de la barre laterale, sauvegarder une fois tout lu
    st.sidebar.subheader("💾 Settings file")
    st.sidebar.file_uploader(
        "Load saved settings (.json)", type=["json"], key="settings_upload",
        on_change=_apply_uploaded_settings,
        help="Colors, hatches, fonts, contribution order and layout saved from a previous session."
    )
    msg = ss.pop("_settings_msg", None)
    if msg:
        getattr(st.sidebar, msg[0])(msg[1])
    settings_box = st.sidebar.container()

    # File uploader
    st.header("Step 2: Upload your filled Excel file")
    uploaded_file = st.file_uploader("Upload your completed Excel file", type=["xlsx"], key="file_upload_main")
    if uploaded_file is None:
        return

    file_bytes = uploaded_file.getvalue()
    data_sig = hashlib.sha1(file_bytes).hexdigest()
    try:
        sheet_names = pd.ExcelFile(io.BytesIO(file_bytes)).sheet_names
    except Exception as e:
        st.error(f"Could not open the Excel file: {e}")
        return

    # Choix de la feuille par l'utilisateur
    selected_sheet = st.selectbox("📄 Choose the sheet to analyze", sheet_names)

    # Affichage aperçu rapide (5 premières lignes)
    with st.expander("👁️ Preview of the selected sheet"):
        st.dataframe(pd.read_excel(io.BytesIO(file_bytes), sheet_name=selected_sheet, nrows=5))

    try:
        initial_table, _combined_table, total_impact_table, scenario_names_cleaned = load_tables(
            data_sig, selected_sheet, file_bytes)
    except TemplateError as e:
        st.error(f"⚠️ The selected sheet does not follow the template. {e}")
        return
    except Exception as e:
        st.error(
            f"⚠️ The selected sheet could not be read ({e}). Please verify that the sheet structure "
            "matches the required format (scenarios in row 1, contributions in row 2, categories in "
            "column A starting from row 3)."
        )
        return

    contributions_list = extract_contributions_from_initial_table(initial_table)

    # Nouveau fichier ou nouvelle feuille : ordre des contributions = ordre du fichier Excel
    order_key = "custom_contribution_order"
    data_key = f"{data_sig}:{selected_sheet}"
    if ss.get("_data_key") != data_key:
        ss["_data_key"] = data_key
        ss[order_key] = list(contributions_list)
    loaded_order = ss.pop("_loaded_order", None)
    if loaded_order:
        valid = [c for c in loaded_order if c in contributions_list]
        if valid:
            ss[order_key] = valid
    ss[order_key] = [c for c in ss.get(order_key, contributions_list) if c in contributions_list]

    st.sidebar.subheader("🔢 Contribution stacking order")
    contributions_order = st.sidebar.multiselect(
        "🪜 Order of contributions (bottom to top)", options=contributions_list, key=order_key
    )
    st.markdown("**Selected contribution order:**")
    for contrib in contributions_order:
        st.markdown(f"- {contrib}")

    font_styles_now = get_font_styles()

    # Couleurs par defaut : palette Okabe-Ito (adaptee aux daltoniens)
    scen_colors = {name.lower(): palette_color(i) for i, name in enumerate(scenario_names_cleaned)}
    contrib_colors = {}
    for contrib in contributions_list:
        contrib_colors.setdefault(contrib.lower(), palette_color(len(contrib_colors)))
    hatches = {name.lower(): DEFAULT_HATCHES[i % len(DEFAULT_HATCHES)]
               for i, name in enumerate(scenario_names_cleaned)}

    # Interface de personnalisation dans la sidebar
    st.sidebar.subheader("🎨 Color customization")
    _ss_default("customize_colors", False)
    if st.sidebar.checkbox("Would you like to customize the colors?", key="customize_colors"):
        st.sidebar.markdown("#### Scenario colors")
        for scen in scenario_names_cleaned:
            key = f"scen_color::{scen.lower()}"
            _ss_default(key, scen_colors[scen.lower()])
            scen_colors[scen.lower()] = st.sidebar.color_picker(f"🎯 {scen.capitalize()}", key=key)

        st.sidebar.markdown("#### Contribution colors")
        for contrib in unique_in_order([c.lower() for c in contributions_list]):
            key = f"contrib_color::{contrib}"
            _ss_default(key, contrib_colors[contrib])
            label = next(c for c in contributions_list if c.lower() == contrib)
            contrib_colors[contrib] = st.sidebar.color_picker(f"📦 {label}", key=key)

    st.sidebar.subheader("🔳 Hatch customization")
    _ss_default("customize_hatches", False)
    if st.sidebar.checkbox("Would you like to customize the hatches?", key="customize_hatches"):
        st.sidebar.markdown("#### Scenario hatches")
        for scen in scenario_names_cleaned:
            scen_lower = scen.lower()
            key = f"hatch_selector_{scen_lower}"
            _ss_default(key, hatches[scen_lower] if hatches[scen_lower] in HATCH_CHOICES else "None")
            hatches[scen_lower] = st.sidebar.selectbox(f"Hatch for {scen.capitalize()}", HATCH_CHOICES, key=key)

    # Choix des graphes
    st.subheader("Select Charts to Generate")
    selected = {}
    for key, label in CHART_CHECKBOXES:
        _ss_default(key, False)
        selected[key] = st.checkbox(label, key=key)
    _ss_default("combined_chart_mode", COMBINED_MODES[0])
    combined_chart_mode = st.selectbox("📊 Combined View of All Scenarios", COMBINED_MODES, key="combined_chart_mode")

    # 📏 Reglage de l'epaisseur des barres et de l'espace entre categories
    combined = None
    if combined_chart_mode != "Do not generate":
        orientation = "vertical" if combined_chart_mode == "Vertical" else "horizontal"
        n_cat = len(initial_table.index)
        n_scen = len(scenario_names_cleaned)
        def_bar_cm, def_gap_cm = default_combined_bar_layout(n_cat, n_scen, orientation)
        def_len_cm = default_combined_bar_length_cm(n_scen, orientation)
        size_word = "width" if orientation == "vertical" else "height"
        length_word = "height" if orientation == "vertical" else "width"

        loaded = ss.get("_loaded_combined") or {}
        use_loaded = loaded.get("orientation") == orientation
        k_bar = f"combined_bar_cm_{orientation}_{n_cat}_{n_scen}"
        k_gap = f"combined_gap_cm_{orientation}_{n_cat}_{n_scen}"
        k_len = f"combined_len_cm_{orientation}_{n_scen}"
        _ss_default(k_bar, _num(loaded.get("bar_cm"), 0.05, 10.0, def_bar_cm) if use_loaded else float(def_bar_cm))
        _ss_default(k_gap, _num(loaded.get("gap_cm"), 0.0, 20.0, def_gap_cm) if use_loaded else float(def_gap_cm))
        _ss_default(k_len, _num(loaded.get("length_cm"), 3.0, 60.0, def_len_cm) if use_loaded else float(def_len_cm))
        _ss_default("combined_show_total_pct", True)
        _ss_default("combined_totals_mode", TOTALS_MODE_LABELS["smart"])
        _ss_default("combined_totals_offset", 6.0)

        with st.expander("📏 Bar layout (combined view)", expanded=True):
            col_bw, col_gap, col_len = st.columns(3)
            bar_cm = col_bw.number_input(
                f"Bar {size_word} (cm)", min_value=0.05, max_value=10.0, step=0.05, key=k_bar,
                help="Thickness of each scenario bar."
            )
            gap_cm = col_gap.number_input(
                "Space between impact categories (cm)", min_value=0.0, max_value=20.0, step=0.05, key=k_gap,
                help="Empty space between two groups of bars (two impact categories)."
            )
            len_cm = col_len.number_input(
                f"Diagram {length_word} (cm)", min_value=3.0, max_value=60.0, step=0.5, key=k_len,
                help="Length of the 0-100 % scale, i.e. how long the bars are. "
                     "Labels and totals are added around it."
            )
            show_pct = st.checkbox(
                "Show percentage in brackets next to totals (Chart with Totals)",
                key="combined_show_total_pct",
                help="Percentage relative to the highest-impact scenario of each category (= 100 %)."
            )
            totals_choice = st.radio(
                "Position of totals (Chart with Totals)", list(TOTALS_MODE_LABELS.values()),
                key="combined_totals_mode",
                help="The axis always stops at 100 %; totals are never crossed by the frame border."
            )
            totals_offset = st.number_input(
                "Gap between bar / frame and totals (pt)", min_value=0.0, max_value=30.0, step=1.0,
                key="combined_totals_offset",
                help="Space between the end of the bar (or the frame border) and the total label. 1 pt ≈ 0.35 mm."
            )
            total_cm = n_cat * (n_scen * bar_cm + gap_cm)
            st.caption(
                f"Resulting diagram (bars area only): {total_cm:.1f} cm {size_word} "
                f"× {len_cm:.1f} cm {length_word}. Defaults reproduce the original layout."
            )
        combined = {
            "orientation": orientation,
            "bar_cm": float(bar_cm),
            "gap_cm": float(gap_cm),
            "length_cm": float(len_cm),
            "show_pct": bool(show_pct),
            "totals_mode": "smart" if totals_choice == TOTALS_MODE_LABELS["smart"] else "row",
            "totals_offset_pt": float(totals_offset),
        }

    # Reglages complets (sauvegarde JSON + ZIP) et reglages utiles au rendu (cle du cache)
    settings = {
        "version": SETTINGS_VERSION,
        "font": font_styles_now,
        "scenario_colors": scen_colors,
        "contribution_colors": contrib_colors,
        "scenario_hatches": hatches,
        "contribution_order": list(contributions_order),
        "charts": {**selected, "combined_chart_mode": combined_chart_mode},
        "combined": combined or {},
    }
    settings_box.download_button(
        "💾 Save current settings (.json)",
        data=json.dumps(settings, indent=2, ensure_ascii=False),
        file_name="lca_graph_settings.json", mime="application/json",
        help="Reload this file later to get exactly the same style (e.g. for all figures of an article)."
    )
    style_json = json.dumps({k: settings[k] for k in
                             ("font", "scenario_colors", "contribution_colors",
                              "scenario_hatches", "contribution_order")}, sort_keys=True)

    jobs = []
    if selected["gen_comparison"]:
        jobs.append(("comparison", "Scenario Comparison", ""))
    if selected["gen_details_v"]:
        jobs.append(("details_v", "Detailed Analysis by Scenario (Vertical)", ""))
    if selected["gen_details_h"]:
        jobs.append(("details_h", "Detailed Analysis by Scenario (Horizontal)", ""))
    if selected["gen_category"]:
        jobs.append(("category", "Detailed Analysis by Impact Category", ""))
    if combined is not None:
        kind = "combined_v" if combined["orientation"] == "vertical" else "combined_h"
        header = "Combined View of All Scenarios" + ("" if kind == "combined_v" else " (Horizontal)")
        jobs.append((kind, header, json.dumps(combined, sort_keys=True)))

    zip_slot = st.container()
    sections = []
    for kind, header, extra_json in jobs:
        st.markdown("---")
        st.header(header)
        with st.spinner("Generating charts…"):
            groups = render_charts(kind, data_sig, selected_sheet, style_json, extra_json, file_bytes)
        for gi, group in enumerate(groups):
            show_group(group, f"{kind}_{gi}")
        sections.append((header, groups))

    if any(groups for _h, groups in sections):
        zip_slot.download_button(
            "🗂️ Download all generated charts (.zip — PNG + SVG + settings)",
            data=build_zip(sections, settings),
            file_name="lca_graphs.zip", mime="application/zip", type="primary",
        )


def main():
    hide_st_style = """
        <style>
        #MainMenu {visibility: hidden;}
        footer {visibility: hidden;}
        </style>
        """
    st.markdown(hide_st_style, unsafe_allow_html=True)

    # 🌐 Barre de navigation
    page = st.radio(" ", ["Introduction", "Graph Generator"], key="page_selector", horizontal=True)

    if page == "Introduction":
        show_introduction_page()
    elif page == "Graph Generator":
        show_generator_page()


if __name__ == "__main__":
    main()
